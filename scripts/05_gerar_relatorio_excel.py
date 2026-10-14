"""
Etapa 5 - Monta o relatorio em Excel (camada de apresentacao para quem nao usa SQL).
Saida: relatorio/relatorio_varejo_online.xlsx
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[1]
RES = RAIZ / "resultados"
ASSETS = RAIZ / "assets"
SAIDA = RAIZ / "relatorio" / "relatorio_varejo_online.xlsx"

FONTE = "Arial"
AZUL_ESCURO = "1F3B5C"
CINZA_CLARO = "F2F2F0"
borda = Border(bottom=Side(style="thin", color="D9D9D6"))

ABAS = [
    ("Qualidade", "01_qualidade_dados.csv", "Qualidade dos dados: o que foi removido e por que"),
    ("KPIs anuais", "02_kpis_anuais.csv", "KPIs anuais (periodos comparaveis de 12 meses)"),
    ("Receita mensal", "03_receita_mensal.csv", "Receita mensal, media movel de 3 meses e variacao anual"),
    ("Novos x recorrentes", "04_novos_vs_recorrentes.csv", "Clientes e receita: novos vs recorrentes"),
    ("Coortes", "05_coortes_retencao.csv", "Retencao por coorte de primeira compra"),
    ("RFM", "06_rfm_segmentacao.csv", "Segmentacao RFM (ultimos 12 meses)"),
    ("Curva ABC", "07_pareto_produtos.csv", "Curva ABC de produtos (ano 2)"),
    ("Cancelamentos", "08_cancelamentos.csv", "Taxa de cancelamento mensal"),
    ("Cancel. produtos", "09_cancelamentos_produtos.csv", "Produtos com maior valor cancelado (ano 2)"),
    ("Paises", "10_paises.csv", "Top 15 paises por receita (ano 2)"),
    ("Dia e hora", "11_dia_hora.csv", "Pedidos por dia da semana e hora"),
    ("Validacao RFM", "12_validacao_rfm.csv", "Validacao do RFM: o que cada segmento de dez/2010 fez nos 12 meses seguintes"),
    ("Dimensionamento", "13_dimensionamento_reativacao.csv", "Quanto vale reativar o segmento 'Em risco' (cenarios)"),
    ("Lista reativacao", "14_lista_reativacao.csv", "Clientes 'Em risco' priorizados para a campanha (maior valor primeiro)"),
]

INSIGHTS = [
    "Receita do ano 2 cresceu 2,5% (bruta) e 1,9% (liquida de cancelamentos), mesmo com 4% menos pedidos.",
    "O crescimento veio do ticket medio (+6,8%), nao de volume nem de novos clientes.",
    "Aquisicao caiu: em set-nov, 600 clientes novos em 2011 contra 940 em 2010 (-36%).",
    "A base ja existente sustenta o negocio: ~90% da receita mensal do ano 2 vem de clientes recorrentes.",
    "Set-nov concentram 37% da receita do ano. Estoque e campanhas precisam estar prontos ate agosto.",
    "RFM: 22% dos clientes (Campeoes) geram 64% da receita. 652 clientes 'Em risco' somam £795 mil.",
    "Curva ABC: 832 de 3.904 produtos (21%) respondem por 80% da receita.",
    "Cancelamento mediano de 2,3% ao mes. O pico de jan/11 (13,7%) e um unico pedido de 74 mil unidades.",
    "Fora do Reino Unido estao 15% da receita. Holanda e Australia tem ticket 5-6x maior (perfil atacado).",
    "Pedidos concentrados de terca a quinta, das 10h as 15h. Quase nenhum pedido aos sabados.",
    "Validacao: segmentando em dez/2010, 94% dos Campeoes voltaram e geraram 68% da receita do ano seguinte.",
    "'Em risco' perde 47% da receita por cliente no ano seguinte. Cada 10 p.p. de reativacao vale £42-80 mil.",
    "103 clientes respondem por metade da receita do segmento 'Em risco': a lista de reativacao comeca por eles.",
]


def cabecalho(ws, linha, colunas):
    for j, nome in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=j, value=nome)
        c.font = Font(name=FONTE, bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[linha].height = 30


def escrever_tabela(ws, df, linha_ini=3):
    cabecalho(ws, linha_ini, list(df.columns))
    for i, row in enumerate(df.itertuples(index=False), start=linha_ini + 1):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=None if pd.isna(v) else v)
            c.font = Font(name=FONTE, size=10)
            c.border = borda
            col = df.columns[j - 1]
            if isinstance(v, (int, float)):
                if "pct" in col or col.startswith("var_") or col == "retencao_pct":
                    c.number_format = '0.0"%"'
                elif col in ("ticket_medio", "valor_medio_cliente", "ticket_vs_reino_unido", "pedidos_por_cliente"):
                    c.number_format = "#,##0.00"
                else:
                    c.number_format = "#,##0"
    for j, col in enumerate(df.columns, start=1):
        largura = max(len(str(col)), *(len(str(x)) for x in df.iloc[:200, j - 1])) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(max(largura, 10), 48)
    ws.freeze_panes = ws.cell(row=linha_ini + 1, column=1)


def titulo(ws, texto):
    ws["A1"] = texto
    ws["A1"].font = Font(name=FONTE, bold=True, size=14, color=AZUL_ESCURO)


def aba_resumo(wb):
    ws = wb.active
    ws.title = "Resumo"
    ws.sheet_view.showGridLines = False
    titulo(ws, "Varejo online (UK), dez/2009 a dez/2011: analise de vendas e clientes em SQL")
    ws["A2"] = "Fonte: Online Retail II, UCI Machine Learning Repository (CC BY 4.0). Valores em libras (£)."
    ws["A2"].font = Font(name=FONTE, italic=True, size=9, color="52514E")

    kpi = pd.read_csv(RES / "02_kpis_anuais.csv")
    linhas = [
        ("Receita bruta (£)", "receita", "#,##0"),
        ("Receita liquida (£)", "receita_liquida", "#,##0"),
        ("Pedidos", "pedidos", "#,##0"),
        ("Clientes identificados", "clientes_identificados", "#,##0"),
        ("Ticket medio (£)", "ticket_medio", "#,##0.00"),
        ("Pedidos por cliente", "pedidos_por_cliente", "0.00"),
    ]
    cabecalho(ws, 4, ["Indicador", "Ano 1 (dez/09-nov/10)", "Ano 2 (dez/10-nov/11)", "Variacao"])
    for i, (nome, col, fmt) in enumerate(linhas, start=5):
        ws.cell(row=i, column=1, value=nome).font = Font(name=FONTE, size=10, bold=True)
        for j in (0, 1):
            c = ws.cell(row=i, column=2 + j, value=float(kpi[col].iloc[j]))
            c.number_format, c.font = fmt, Font(name=FONTE, size=10)
        v = ws.cell(row=i, column=4, value=f"=C{i}/B{i}-1")
        v.number_format, v.font = "+0.0%;-0.0%", Font(name=FONTE, size=10, bold=True)
        for j in range(1, 5):
            ws.cell(row=i, column=j).border = borda

    ws["A13"] = "Principais conclusoes"
    ws["A13"].font = Font(name=FONTE, bold=True, size=12, color=AZUL_ESCURO)
    for i, txt in enumerate(INSIGHTS, start=14):
        c = ws.cell(row=i, column=1, value=f"{i - 13}. {txt}")
        c.font = Font(name=FONTE, size=10)
    ws.column_dimensions["A"].width = 30
    for col in "BCD":
        ws.column_dimensions[col].width = 22


def grafico_receita(ws, n):
    ch = LineChart()
    ch.title = "Receita mensal (£)"
    ch.height, ch.width = 8, 22
    ch.add_data(Reference(ws, min_col=2, min_row=3, max_row=3 + n), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=5, min_row=3, max_row=3 + n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + n))
    ch.series[0].graphicalProperties.line.solidFill = "2A78D6"
    ch.series[1].graphicalProperties.line.solidFill = "A3A29C"
    ch.series[1].graphicalProperties.line.dashStyle = "dash"
    ws.add_chart(ch, "J3")


def grafico_rfm(ws, n):
    ch = BarChart()
    ch.type = "bar"
    ch.title = "RFM: % clientes vs % receita"
    ch.height, ch.width = 8, 18
    ch.add_data(Reference(ws, min_col=3, min_row=3, max_row=3 + n), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=5, min_row=3, max_row=3 + n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=3 + n))
    ch.series[0].graphicalProperties.solidFill = "A3A29C"
    ch.series[1].graphicalProperties.solidFill = "2A78D6"
    ws.add_chart(ch, "J3")


def aba_graficos(wb):
    ws = wb.create_sheet("Graficos")
    ws.sheet_view.showGridLines = False
    titulo(ws, "Graficos")
    linha = 3
    for png in sorted(ASSETS.glob("*.png")):
        img = XLImage(str(png))
        escala = 900 / img.width
        img.width, img.height = int(img.width * escala), int(img.height * escala)
        ws.add_image(img, f"A{linha}")
        linha += int(img.height / 20) + 2


def main():
    SAIDA.parent.mkdir(exist_ok=True)
    wb = Workbook()
    aba_resumo(wb)
    aba_graficos(wb)
    for nome, arquivo, texto in ABAS:
        df = pd.read_csv(RES / arquivo)
        ws = wb.create_sheet(nome)
        titulo(ws, texto)
        escrever_tabela(ws, df)
        if nome == "Receita mensal":
            grafico_receita(ws, len(df))
        if nome == "RFM":
            grafico_rfm(ws, len(df))
    wb.save(SAIDA)
    print("ok", SAIDA.name)


if __name__ == "__main__":
    main()
